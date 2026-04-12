"""
ContactAudit.py — GlueUp Contact Audit Tool
ICF Washington State Chapter

Cross-references a GlueUp Admin UI contact export (XLSX) against the live
GlueUp membership directory API to identify:

  1. INACTIVE CONTACTS — contacts with no active ICF Chapter Member (37600)
     membership. These may be unlinked, or legitimately non-members.

  2. DUPLICATE NAMES — contacts with the same name but different email
     addresses. High risk of being the same person (Case 3 — email mismatch
     between GlueUp and ICF Global).

  3. SHADOW EMAIL CONTACTS — contacts using id_*@members.icfwashingtonstate.org
     proxy emails. Flagged so you can check if a real email is now available.

  4. ACTIVE MEMBER / API MISMATCH — contacts marked Active Member=Yes in the
     export but not found in the live API membership directory (or vice versa).

Loads the session token from glueup_token.json (run GetToken.py first).

Usage:
    python3 ContactAudit.py --in contacts_export.xlsx
    python3 ContactAudit.py --in contacts_export.xlsx --out my_audit.xlsx
    python3 ContactAudit.py --in contacts_export.xlsx --csv   # output as CSV instead

Notes on the input file format:
    - Export from GlueUp Admin UI: Contacts → Export
    - Row 1: column headers
    - Row 2: Volunteer Role sub-headers (Beta Tester, Pro Bono Coach, etc.)
    - Data starts at Row 3
"""

import argparse
import json
import hmac
import hashlib
import time
import os
import sys
import re
from collections import defaultdict
from datetime import datetime
from urllib.request import urlopen, Request
from urllib.error import HTTPError

try:
    from openpyxl import load_workbook
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment
except ImportError:
    print('✗ openpyxl not installed. Run: pip3 install openpyxl --break-system-packages')
    sys.exit(1)

# ── GlueUp API constants ──────────────────────────────────────────────────────
GLUEUP_BASE_URL    = 'https://api-services.glueup.com/v2'
ORG_ID             = '7912'
PUBLIC_KEY         = 'icfwshts'
PRIVATE_KEY        = ('MF4CAQACEADAzyLnyJdLTnvVextU0XMCAwEAAQIQAKT41snxxRoPf'
                      'Xb0gguT2QIIDoAWHftn8h0CCA1L/8fNanzPAggAksHNF6ZpZQIIARa17LAfBfU'
                      'CCAP67vZiAQ55')
ICF_MEMBERSHIP_TYPE = 37600
SHADOW_PATTERN      = re.compile(r'^id_\d+@members\.icfwashingtonstate\.org$', re.IGNORECASE)
TOKEN_FILE          = os.path.join(os.path.dirname(__file__), 'glueup_token.json')

# Colours for Excel output
FILL_HEADER   = PatternFill('solid', fgColor='1F4E79')  # dark blue
FILL_SECTION  = PatternFill('solid', fgColor='2E75B6')  # medium blue
FILL_WARN     = PatternFill('solid', fgColor='FFE699')  # yellow
FILL_RISK     = PatternFill('solid', fgColor='F4CCCC')  # red-pink
FILL_OK       = PatternFill('solid', fgColor='D9EAD3')  # green
FONT_WHITE    = Font(bold=True, color='FFFFFF')
FONT_BOLD     = Font(bold=True)


# ── Auth helpers ──────────────────────────────────────────────────────────────

def build_a_header(method: str) -> str:
    ts     = str(int(time.time() * 1000))
    msg    = method + PUBLIC_KEY + '1.0' + ts
    digest = hmac.new(PRIVATE_KEY.encode(), msg.encode(), hashlib.sha256).hexdigest()
    return f'v=1.0;k={PUBLIC_KEY};ts={ts};d={digest}'


def load_token() -> str:
    if not os.path.exists(TOKEN_FILE):
        print('✗ No token file found. Run GetToken.py first.')
        sys.exit(1)
    with open(TOKEN_FILE) as f:
        data = json.load(f)
    return data['token']


def glueup_post(endpoint: str, token: str, body: dict) -> dict:
    a_header = build_a_header('POST')
    encoded  = json.dumps(body).encode()
    req = Request(
        f'{GLUEUP_BASE_URL}{endpoint}',
        data=encoded,
        headers={
            'Content-Type':          'application/json',
            'a':                     a_header,
            'token':                 token,
            'requestOrganizationId': ORG_ID,
            'User-Agent':            'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7)'
        },
        method='POST'
    )
    with urlopen(req) as resp:
        return json.loads(resp.read())


# ── API fetch ─────────────────────────────────────────────────────────────────

def fetch_all_members(token: str) -> list:
    """Fetch all active members from the GlueUp membership directory."""
    all_records = []
    page_size   = 50
    offset      = 0

    print('Fetching live membership data from GlueUp API...', end='', flush=True)

    while True:
        body = {
            'projection': [],
            'offset':     offset,
            'limit':      page_size,
            'order':      {'familyName': 'asc'}
        }
        result  = glueup_post('/membershipDirectory/members', token, body)
        records = result.get('value', [])

        if not records:
            break

        all_records.extend(records)
        print('.', end='', flush=True)

        if len(records) < page_size:
            break

        offset += page_size

    print(f' done. ({len(all_records)} active membership records)\n')
    return all_records


def build_api_index(api_records: list) -> tuple[set, dict]:
    """
    Build lookup structures from the API response.
    Returns:
        icf_member_emails  — set of emails with an active 37600 membership
        api_by_email       — dict of email → full API record (for ICF field enrichment)
    """
    icf_member_emails = set()
    api_by_email      = {}

    for rec in api_records:
        member     = rec.get('individualMember', {})
        membership = rec.get('membership', {})
        email      = (member.get('emailAddress') or {}).get('value', '').lower().strip()

        if not email:
            continue

        api_by_email[email] = rec

        mt_id = (membership.get('membershipType') or {}).get('id')
        if mt_id == ICF_MEMBERSHIP_TYPE:
            icf_member_emails.add(email)

    return icf_member_emails, api_by_email


# ── Export file reader ────────────────────────────────────────────────────────

def read_export(path: str) -> list[dict]:
    """
    Read the GlueUp Admin contact export XLSX.

    Row 1: column headers
    Row 2: Volunteer Role sub-headers (Beta Tester, Pro Bono Coach, Board Member, Member Support)
    Row 3+: data

    Volunteer Role is a multi-column checkbox export — columns 7-10 (0-indexed 6-9).
    Each is None or 'x'. We consolidate them into a pipe-separated string.
    """
    wb = load_workbook(path, read_only=True)
    ws = wb.active

    rows     = list(ws.iter_rows(values_only=True))
    headers  = rows[0]   # Row 1
    vol_hdrs = rows[1]   # Row 2 — volunteer role sub-headers

    # Find volunteer role column range: contiguous None-headed columns after col 6
    # In the sample: cols 6,7,8,9 (0-indexed) are the 4 volunteer role options
    vol_col_names = []
    vol_col_idxs  = []
    for i, h in enumerate(headers):
        if h == 'Volunteer Role':
            vol_col_idxs.append(i)
            vol_col_names.append(vol_hdrs[i] or f'VolRole_{i}')
        elif h is None and i > 0 and headers[i-1] in (None, 'Volunteer Role'):
            # continuation of the multi-column block
            if vol_col_idxs and i == vol_col_idxs[-1] + 1:
                vol_col_idxs.append(i)
                vol_col_names.append(vol_hdrs[i] or f'VolRole_{i}')

    contacts = []
    for row in rows[2:]:  # data starts at row 3
        if not any(row):
            continue

        email = str(row[3] or '').strip().lower()
        if not email:
            continue

        # Consolidate volunteer roles
        vol_roles = []
        for idx, name in zip(vol_col_idxs, vol_col_names):
            if idx < len(row) and row[idx] == 'x':
                vol_roles.append(name)

        # Date Created — may be a datetime object or string
        date_created = row[25]
        if isinstance(date_created, datetime):
            date_created = date_created.strftime('%Y-%m-%d')
        elif date_created:
            date_created = str(date_created)
        else:
            date_created = ''

        contacts.append({
            'firstName':     str(row[1] or '').strip(),
            'lastName':      str(row[2] or '').strip(),
            'email':         email,
            'company':       str(row[4] or '').strip(),
            'title':         str(row[5] or '').strip(),
            'volunteerRoles': '|'.join(vol_roles),
            'findable':      str(row[10] or '').strip(),
            'contactId':     row[14],
            'activeMember':  str(row[15] or '').strip(),  # Yes / No
            'dateCreated':   date_created,
            'isShadow':      bool(SHADOW_PATTERN.match(email)),
        })

    return contacts


# ── Audit logic ───────────────────────────────────────────────────────────────

def run_audit(contacts: list, icf_member_emails: set, api_by_email: dict) -> dict:
    """
    Run all four audit checks and return categorised results.
    """
    # 1 — Inactive contacts (not in icf_member_emails)
    inactive = [c for c in contacts if c['email'] not in icf_member_emails]

    # 2 — Duplicate names (same normalised name, different emails)
    by_name = defaultdict(list)
    for c in contacts:
        name = f'{c["firstName"]} {c["lastName"]}'.strip().lower()
        # Normalise titles/credentials in names (strip everything after first comma)
        name = name.split(',')[0].strip()
        if name:
            by_name[name].append(c)
    dup_names = {n: contacts for n, contacts in by_name.items() if len(contacts) > 1}

    # 3 — Shadow email contacts
    shadow = [c for c in contacts if c['isShadow']]

    # 4 — Active Member / API mismatch
    #   a) Export says Active Member=Yes but not in API icf_member_emails
    #   b) In API icf_member_emails but export says Active Member=No
    export_emails = {c['email'] for c in contacts}
    mismatch_export_yes_api_no = [
        c for c in contacts
        if c['activeMember'] == 'Yes' and c['email'] not in icf_member_emails
    ]
    mismatch_api_yes_export_no = [
        email for email in icf_member_emails
        if email in export_emails and
        next((c for c in contacts if c['email'] == email), {}).get('activeMember') != 'Yes'
    ]

    # Enrich inactive contacts with ICF fields from API where available
    for c in inactive:
        api_rec = api_by_email.get(c['email'])
        if api_rec:
            member = api_rec.get('individualMember', {})
            props  = member.get('properties') or {}
            c['icfmemberid']   = props.get('icfmemberid', '')
            c['icfimportdate'] = props.get('icfimportdate', '')
            c['membershipType'] = (api_rec.get('membership', {})
                                   .get('membershipType', {})
                                   .get('title', ''))
        else:
            c['icfmemberid']    = ''
            c['icfimportdate']  = ''
            c['membershipType'] = ''

    return {
        'inactive':                   inactive,
        'duplicate_names':            dup_names,
        'shadow':                     shadow,
        'mismatch_export_yes_api_no': mismatch_export_yes_api_no,
        'mismatch_api_yes_export_no': mismatch_api_yes_export_no,
    }


# ── Output — console summary ──────────────────────────────────────────────────

def print_summary(results: dict, total: int) -> None:
    inactive   = results['inactive']
    dup_names  = results['duplicate_names']
    shadow     = results['shadow']
    mm_yn      = results['mismatch_export_yes_api_no']
    mm_ny      = results['mismatch_api_yes_export_no']

    print('=' * 72)
    print(f'CONTACT AUDIT SUMMARY  ({total} total contacts in export)')
    print('=' * 72)
    print(f'  1. Contacts without ICF Chapter membership:  {len(inactive)}')
    print(f'  2. Duplicate name groups (diff emails):      {len(dup_names)}')
    print(f'  3. Shadow email contacts:                    {len(shadow)}')
    print(f'  4a. Export=Active but not in API directory:  {len(mm_yn)}')
    print(f'  4b. In API directory but Export=Inactive:    {len(mm_ny)}')
    print('=' * 72)

    if dup_names:
        print(f'\n── Duplicate Names ──')
        for name, contacts in dup_names.items():
            print(f'\n  {name.title()}')
            for c in contacts:
                shadow_flag = ' [SHADOW]' if c['isShadow'] else ''
                print(f'    contactId={c["contactId"]}  '
                      f'email={c["email"]}{shadow_flag}  '
                      f'activeMember={c["activeMember"]}')

    if shadow:
        print(f'\n── Shadow Email Contacts ──')
        for c in shadow:
            print(f'  {c["firstName"]} {c["lastName"]}  '
                  f'email={c["email"]}  activeMember={c["activeMember"]}')

    if mm_yn:
        print(f'\n── Export says Active but not found in API (type 37600) ──')
        for c in mm_yn:
            print(f'  {c["firstName"]} {c["lastName"]}  email={c["email"]}')

    if mm_ny:
        print(f'\n── In API directory (type 37600) but Export says Inactive ──')
        for email in mm_ny:
            print(f'  {email}')


# ── Output — Excel workbook ───────────────────────────────────────────────────

def write_excel(results: dict, total: int, out_path: str) -> None:
    wb = Workbook()

    def make_sheet(title: str) -> object:
        ws = wb.create_sheet(title=title)
        return ws

    def hdr(ws, values: list, fill=FILL_HEADER, font=FONT_WHITE) -> None:
        ws.append(values)
        for cell in ws[ws.max_row]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(wrap_text=True)

    def row_fill(ws, fill) -> None:
        for cell in ws[ws.max_row]:
            cell.fill = fill

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = 'Summary'
    hdr(ws_sum, ['Audit Check', 'Count', 'Action Required'])
    checks = [
        ('Contacts without ICF Chapter membership (37600)',
         len(results['inactive']),
         'Review — may be unlinked or legitimately non-members'),
        ('Duplicate name groups (different email addresses)',
         len(results['duplicate_names']),
         'HIGH RISK — likely same person with email mismatch'),
        ('Shadow email contacts',
         len(results['shadow']),
         'Check if real email is now available; promote if so'),
        ('Export=Active but missing from API directory',
         len(results['mismatch_export_yes_api_no']),
         'Investigate — possible stale export or API sync issue'),
        ('In API directory but Export shows Inactive',
         len(results['mismatch_api_yes_export_no']),
         'Investigate — possible export filter or timing issue'),
        ('', '', ''),
        ('Total contacts in export', total, ''),
        ('Audit run date', datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC'), ''),
    ]
    for check, count, action in checks:
        ws_sum.append([check, count, action])
        if isinstance(count, int) and count > 0:
            row_fill(ws_sum, FILL_WARN)
    ws_sum.column_dimensions['A'].width = 55
    ws_sum.column_dimensions['B'].width = 10
    ws_sum.column_dimensions['C'].width = 55

    # ── Sheet 2: Duplicate Names ──────────────────────────────────────────────
    ws_dup = make_sheet('Duplicate Names')
    hdr(ws_dup, ['Name', 'Contact ID', 'Email', 'Active Member',
                 'Shadow Email', 'Volunteer Roles', 'Date Created', 'Notes'])
    for name, contacts in results['duplicate_names'].items():
        for i, c in enumerate(contacts):
            notes = 'SHADOW EMAIL' if c['isShadow'] else ''
            ws_dup.append([
                name.title() if i == 0 else '',
                c['contactId'],
                c['email'],
                c['activeMember'],
                'Yes' if c['isShadow'] else 'No',
                c['volunteerRoles'],
                c['dateCreated'],
                notes
            ])
            row_fill(ws_dup, FILL_RISK if c['isShadow'] else FILL_WARN)
    for col, width in zip('ABCDEFGH', [30, 12, 40, 14, 14, 30, 20, 20]):
        ws_dup.column_dimensions[col].width = width

    # ── Sheet 3: Shadow Emails ────────────────────────────────────────────────
    ws_shad = make_sheet('Shadow Emails')
    hdr(ws_shad, ['First Name', 'Last Name', 'Shadow Email', 'ICF Member ID',
                  'Contact ID', 'Active Member', 'Date Created'])
    for c in results['shadow']:
        icf_id = c.get('icfmemberid', '')
        # Extract ICF member ID from shadow email as fallback
        if not icf_id:
            m = re.match(r'^id_(\d+)@', c['email'])
            icf_id = m.group(1) if m else ''
        ws_shad.append([
            c['firstName'], c['lastName'], c['email'],
            icf_id, c['contactId'], c['activeMember'], c['dateCreated']
        ])
        row_fill(ws_shad, FILL_WARN)
    for col, width in zip('ABCDEFG', [15, 15, 45, 15, 12, 14, 20]):
        ws_shad.column_dimensions[col].width = width

    # ── Sheet 4: Inactive Contacts ────────────────────────────────────────────
    ws_inact = make_sheet('No ICF Membership')
    hdr(ws_inact, ['First Name', 'Last Name', 'Email', 'Shadow Email',
                   'Active Member', 'Other Membership', 'ICF Member ID',
                   'ICF Import Date', 'Company', 'Volunteer Roles',
                   'Contact ID', 'Date Created'])
    for c in results['inactive']:
        ws_inact.append([
            c['firstName'], c['lastName'], c['email'],
            'Yes' if c['isShadow'] else 'No',
            c['activeMember'],
            c.get('membershipType', ''),
            c.get('icfmemberid', ''),
            c.get('icfimportdate', ''),
            c['company'],
            c['volunteerRoles'],
            c['contactId'],
            c['dateCreated'],
        ])
        fill = FILL_RISK if c['isShadow'] else (FILL_OK if c['activeMember'] == 'No' else FILL_WARN)
        row_fill(ws_inact, fill)
    for col, width in zip('ABCDEFGHIJKL', [15, 15, 40, 12, 14, 20, 15, 15, 25, 30, 12, 20]):
        ws_inact.column_dimensions[col].width = width

    # ── Sheet 5: Mismatches ───────────────────────────────────────────────────
    ws_mm = make_sheet('Mismatches')
    hdr(ws_mm, ['Issue', 'First Name', 'Last Name', 'Email', 'Export Active Member'])
    for c in results['mismatch_export_yes_api_no']:
        ws_mm.append(['Export=Active, not in API', c['firstName'], c['lastName'],
                      c['email'], c['activeMember']])
        row_fill(ws_mm, FILL_RISK)
    for email in results['mismatch_api_yes_export_no']:
        ws_mm.append(['In API, Export=Inactive', '', '', email, 'No'])
        row_fill(ws_mm, FILL_WARN)
    for col, width in zip('ABCDE', [35, 15, 15, 40, 18]):
        ws_mm.column_dimensions[col].width = width

    wb.save(out_path)
    print(f'\nAudit workbook saved to: {out_path}')


# ── Output — CSV (flat, one sheet = No ICF Membership) ───────────────────────

def write_csv(results: dict, out_path: str) -> None:
    import csv
    fieldnames = ['firstName', 'lastName', 'email', 'isShadow', 'activeMember',
                  'membershipType', 'icfmemberid', 'icfimportdate',
                  'company', 'volunteerRoles', 'contactId', 'dateCreated']
    with open(out_path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction='ignore')
        writer.writeheader()
        for c in results['inactive']:
            c['isShadow'] = 'Yes' if c['isShadow'] else 'No'
            writer.writerow(c)
    print(f'CSV saved to: {out_path}')


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description='Audit GlueUp contacts: duplicates, shadow emails, unlinked members.'
    )
    parser.add_argument('--in',  dest='infile', required=True,
                        help='GlueUp Admin contact export XLSX file')
    parser.add_argument('--out', dest='outfile', default=None,
                        help='Output file path (default: contact_audit_YYYYMMDD.xlsx)')
    parser.add_argument('--csv', action='store_true',
                        help='Output as CSV instead of Excel (No ICF Membership sheet only)')
    args = parser.parse_args()

    if not os.path.exists(args.infile):
        print(f'✗ Input file not found: {args.infile}')
        sys.exit(1)

    # Default output filename
    if not args.outfile:
        date_str   = datetime.utcnow().strftime('%Y%m%d')
        ext        = '.csv' if args.csv else '.xlsx'
        args.outfile = os.path.join(os.path.dirname(args.infile),
                                    f'contact_audit_{date_str}{ext}')

    print(f'Reading contact export: {args.infile}')
    contacts = read_export(args.infile)
    print(f'Loaded {len(contacts)} contacts from export.\n')

    token                      = load_token()
    api_records                = fetch_all_members(token)
    icf_member_emails, api_by_email = build_api_index(api_records)

    results = run_audit(contacts, icf_member_emails, api_by_email)
    print_summary(results, len(contacts))

    if args.csv:
        write_csv(results, args.outfile)
    else:
        write_excel(results, len(contacts), args.outfile)


if __name__ == '__main__':
    main()